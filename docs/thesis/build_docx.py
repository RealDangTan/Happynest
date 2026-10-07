"""Build the thesis DOCX from the chapter Markdown files.

Run with the bundled document Python runtime. The source Markdown remains the
editable, agent-owned version; the DOCX is the assembled delivery artifact.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Khoa-luan-Le-Dang-Tan-11236199.docx"
BUILD = ROOT / ".build"
BUILD.mkdir(exist_ok=True)
TITLE = "AI AGENT TỔNG HỢP, PHÂN LOẠI VÀ PHÁT HIỆN VẤN ĐỀ TỪ PHẢN HỒI NGƯỜI DÙNG VỀ SẢN PHẨM AI"
FILES = [
    "02-chuong-1-tong-quan.md",
    "03-chuong-2-co-so-ly-thuyet.md",
    "04-chuong-3-phan-tich-thiet-ke.md",
    "05-chuong-4-ket-qua-danh-gia.md",
    "06-ket-luan.md",
    "07-tai-lieu-tham-khao.md",
    "09-phu-luc.md",
]


def xml_child(parent, name, attrs=None, text=None):
    child = OxmlElement(name)
    for key, value in (attrs or {}).items():
        child.set(qn(key), str(value))
    if text is not None:
        child.text = text
    parent.append(child)
    return child


def field(paragraph, instruction, result=""):
    r = paragraph.add_run()
    xml_child(r._r, "w:fldChar", {"w:fldCharType": "begin"})
    r = paragraph.add_run()
    xml_child(r._r, "w:instrText", {"xml:space": "preserve"}, f" {instruction} ")
    r = paragraph.add_run()
    xml_child(r._r, "w:fldChar", {"w:fldCharType": "separate"})
    if result:
        paragraph.add_run(result)
    r = paragraph.add_run()
    xml_child(r._r, "w:fldChar", {"w:fldCharType": "end"})


def page_number(section, roman=False):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, "PAGE \\* ROMAN" if roman else "PAGE", "i" if roman else "1")


def set_page_number_start(section, start=1, roman=False):
    sect = section._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    xml_child(sect, "w:pgNumType", {"w:start": str(start), "w:fmt": "lowerRoman" if roman else "decimal"})


def set_margins(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.3)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(18 if level == 1 else 13)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.space_before = Pt(6 if level > 1 else 0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.3
        pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT


def add_inline(paragraph, text):
    # Minimal Markdown inline syntax needed by the source manuscript.
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1]); run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def add_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    return p


def add_heading(doc, text, level, new_page=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if new_page:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return p


def make_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(row[j] if j < len(row) else "")
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.bold = i == 0
            if i == 0:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7EDF4")
                tcpr.append(shd)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        xml_child(trpr, "w:cantSplit")
    doc.add_paragraph()


def font(size=28, bold=False):
    path = Path("C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf")
    if path.exists():
        return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def make_diagram(index):
    # Diagrams are deliberately simple vector-like drawings derived from the
    # manuscript's three Mermaid descriptions; no empirical content is added.
    diagrams = {
        1: ["Tệp CSV / nhập tay", "LISTEN: hồ sơ cột", "Gate 1: xác nhận", "Che PII & lưu", "Phân loại & vector", "UNDERSTAND", "Gate 2: duyệt", "ACT", "Gate 3: duyệt", "Theo dõi"],
        2: ["Người dùng", "Next.js", "FastAPI API", "Dịch vụ LISTEN / ACT", "LangGraph UNDERSTAND", "PostgreSQL + pgvector", "LLM / embedding", "Langfuse / logs"],
        3: ["load_context", "planner", "dispatch công cụ", "evaluator", "synthesizer", "persist_insight", "Gate 2: interrupt", "apply_decision", "Kết thúc / điều tra tiếp"],
    }
    labels = diagrams[index]
    width = 1500
    ncols = 2
    rows = (len(labels) + 1) // 2
    cellw, cellh = 650, 100
    height = rows * 150 + 80
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f = font(28, True)
    coords = []
    for i, label in enumerate(labels):
        col = i % 2
        row = i // 2
        x0, y0 = 50 + col * 750, 35 + row * 150
        x1, y1 = x0 + cellw, y0 + cellh
        draw.rounded_rectangle((x0, y0, x1, y1), radius=16, fill="#EAF1F8", outline="#31567D", width=3)
        box = draw.textbbox((0, 0), label, font=f)
        draw.text(((x0+x1-(box[2]-box[0]))/2, (y0+y1-(box[3]-box[1]))/2-3), label, font=f, fill="#183B5C")
        coords.append((x0, y0, x1, y1))
    for i in range(len(coords)-1):
        a, b = coords[i], coords[i+1]
        if i % 2 == 0:
            start, end = (a[2]+5, (a[1]+a[3])//2), (b[0]-7, (b[1]+b[3])//2)
        else:
            start, end = ((a[0]+a[2])//2, a[3]+4), ((b[0]+b[2])//2, b[1]-7)
        draw.line((start, end), fill="#31567D", width=4)
        ex, ey = end
        if i % 2 == 0:
            draw.polygon([(ex,ey),(ex-16,ey-10),(ex-16,ey+10)], fill="#31567D")
        else:
            draw.polygon([(ex,ey),(ex-10,ey-16),(ex+10,ey-16)], fill="#31567D")
    out = BUILD / f"hinh-3-{index}.png"
    image.save(out)
    return out


def add_caption(doc, line):
    match = re.match(r"\*\*(Bảng|Hình)\s+(\d+)\.(\d+)\.\s*(.*?)\*\*$", line)
    if not match:
        return False
    kind, chapter, number, description = match.groups()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{kind} {chapter}.")
    run.bold = True
    field(p, f"SEQ {kind} \\* ARABIC \\s 1", number)
    p.add_run(f". {description}").bold = True
    return True


def parse_markdown(doc, path, first_page=False):
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    i = 0
    current_caption = ""
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("```mermaid"):
            match = re.search(r"Hình 3\.(\d+)", current_caption)
            if match:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(make_diagram(int(match.group(1)))), width=Cm(14.5))
                add_caption(doc, current_caption)
            while i < len(lines) and not lines[i].strip() == "```":
                i += 1
            i += 1; current_caption = ""; continue
        if line.startswith("```"):
            i += 1
            while i < len(lines) and lines[i].strip() != "```":
                p = add_paragraph(doc, lines[i]); p.paragraph_format.left_indent = Cm(0.5)
                i += 1
            i += 1; continue
        image_match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
        if image_match:
            image_path = (ROOT / image_match.group(1)).resolve()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image_path), width=Cm(14.5))
            if current_caption:
                add_caption(doc, current_caption)
                current_caption = ""
            i += 1; continue
        formula_match = re.match(r"\*\*\((\d+)\.(\d+)\)\*\*\s*`([^`]+)`", line)
        if formula_match:
            chapter, number, expression = formula_match.groups()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.tab_stops.add_tab_stop(Cm(15), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(expression)
            p.add_run("\t(" + chapter + ".")
            field(p, "SEQ Công thức \\* ARABIC \\s 1", number)
            p.add_run(")")
            i += 1; continue
        if line.startswith("# "):
            add_heading(doc, line[2:], 1, new_page=first_page)
            first_page = False; i += 1; continue
        if line.startswith("## "):
            add_heading(doc, line[3:], 2); i += 1; continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 3); i += 1; continue
        if re.match(r"\*\*Hình\s+\d+\.\d+\.", line):
            current_caption = line; i += 1; continue
        if add_caption(doc, line):
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip().replace("`", "") for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            make_table(doc, rows); continue
        if re.match(r"^[-*] ", line):
            add_paragraph(doc, line[2:], "List Bullet")
            i += 1; continue
        if re.match(r"^\d+\. ", line):
            add_paragraph(doc, re.sub(r"^\d+\.\s+", "", line), "List Number")
            i += 1; continue
        add_paragraph(doc, line)
        i += 1


def cover(doc, secondary=False):
    section = doc.sections[-1]
    section.different_first_page_header_footer = True
    center = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = center
    p.add_run("ĐẠI HỌC KINH TẾ QUỐC DÂN").bold = True
    p = doc.add_paragraph(); p.alignment = center
    p.add_run("TRƯỜNG CÔNG NGHỆ – KHOA CÔNG NGHỆ THÔNG TIN").bold = True
    doc.add_paragraph()
    with ZipFile(ROOT / "sample_file.docx") as z:
        logo = io.BytesIO(z.read("word/media/image1.jpeg"))
    p = doc.add_paragraph(); p.alignment = center
    p.add_run().add_picture(logo, width=Cm(3.3))
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = center
    r = p.add_run("KHÓA LUẬN TỐT NGHIỆP"); r.bold = True; r.font.size = Pt(20)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = center
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(16)
    for _ in range(3): doc.add_paragraph()
    for label, value in [
        ("Sinh viên thực hiện", "Lê Đăng Tấn"),
        ("Mã sinh viên", "11236199"),
        ("Lớp", "CNTT65A"),
        ("Ngành", "Công nghệ thông tin"),
        ("Giảng viên hướng dẫn", "TS. Lê Thị Hoài Thu"),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(2.1)
        p.add_run(label + ": ").bold = True
        p.add_run(value)
    if secondary:
        doc.add_paragraph()
        p = doc.add_paragraph(); p.alignment = center
        p.add_run("Xác nhận của giảng viên hướng dẫn").bold = True
        doc.add_paragraph(); doc.add_paragraph()
        p = doc.add_paragraph(); p.alignment = center
        p.add_run("...................................................")
    for _ in range(2 if secondary else 4): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = center
    p.add_run("Hà Nội, tháng 12 năm 2026")


def simple_front_section(doc, title, body):
    add_heading(doc, title, 1, new_page=True)
    for paragraph in body.split("\n\n"):
        if paragraph.strip():
            add_paragraph(doc, paragraph.strip())


def frontmatter(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.add_run("MỤC LỤC")
    p = doc.add_paragraph()
    field(p, 'TOC \\o "1-3" \\h \\z \\u', "Nhấn Ctrl+A, F9 trong Word để cập nhật mục lục tự động.")
    front = (ROOT / "01-phan-dau.md").read_text(encoding="utf-8")
    declaration = front.split("## Lời cam đoan", 1)[1].split("## Lời cảm ơn", 1)[0].strip()
    thanks = front.split("## Lời cảm ơn", 1)[1].split("## Danh mục từ viết tắt", 1)[0].strip()
    simple_front_section(doc, "LỜI CAM ĐOAN", declaration)
    simple_front_section(doc, "LỜI CẢM ƠN", thanks)
    add_heading(doc, "DANH MỤC TỪ VIẾT TẮT", 1, new_page=True)
    part = front.split("## Danh mục từ viết tắt", 1)[1].split("## MỞ ĐẦU", 1)[0]
    rows = []
    for line in part.splitlines():
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells): rows.append(cells)
    make_table(doc, rows)
    for title, kind in [("DANH MỤC BẢNG VÀ BIỂU ĐỒ", "Bảng"), ("DANH MỤC HÌNH", "Hình"), ("DANH MỤC CÔNG THỨC", "Công thức")]:
        add_heading(doc, title, 1, new_page=True)
        p = doc.add_paragraph()
        field(p, f'TOC \\h \\z \\c "{kind}"', "Nhấn Ctrl+A, F9 trong Word để cập nhật danh mục tự động.")
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_margins(section); set_page_number_start(section, 1); page_number(section)
    add_heading(doc, "MỞ ĐẦU", 1)
    opening = front.split("## MỞ ĐẦU", 1)[1].strip()
    for paragraph in opening.split("\n\n"):
        if paragraph.strip(): add_paragraph(doc, paragraph.strip())


def main():
    doc = Document()
    configure_styles(doc)
    set_margins(doc.sections[0])
    cover(doc)
    doc.add_page_break()
    cover(doc, secondary=True)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_margins(section); set_page_number_start(section, 1, roman=True); page_number(section, roman=True)
    frontmatter(doc)
    for path in FILES:
        parse_markdown(doc, path, first_page=True)
    add_heading(doc, "KẾT QUẢ KIỂM TRA TURNITIN", 1, new_page=True)
    add_paragraph(doc, "Chưa có kết quả Turnitin tại thời điểm soạn bản DOCX. Cần chèn phiếu kết quả chính thức trước khi nộp bản cuối.")
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
