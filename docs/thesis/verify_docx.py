"""Structural QA for the assembled thesis document (not a visual Word render)."""

from pathlib import Path
from zipfile import ZipFile

from docx import Document


path = Path(__file__).resolve().parent / "Khoa-luan-Le-Dang-Tan-11236199.docx"
with ZipFile(path) as archive:
    assert archive.testzip() is None, "DOCX ZIP has a corrupt member"
    xml = archive.read("word/document.xml").decode("utf-8")
    media = [name for name in archive.namelist() if name.startswith("word/media/")]

doc = Document(path)
headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
required = [
    "MỤC LỤC", "LỜI CAM ĐOAN", "LỜI CẢM ƠN", "DANH MỤC TỪ VIẾT TẮT",
    "DANH MỤC BẢNG VÀ BIỂU ĐỒ", "DANH MỤC HÌNH", "DANH MỤC CÔNG THỨC",
    "MỞ ĐẦU", "CHƯƠNG 1.", "CHƯƠNG 2.", "CHƯƠNG 3.", "CHƯƠNG 4.",
    "KẾT LUẬN", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC", "KẾT QUẢ KIỂM TRA TURNITIN",
]
assert len(headings) == len(required), (len(headings), headings)
for item, heading in zip(required, headings):
    assert heading.startswith(item), (item, heading)
all_text = "\n".join(p.text for p in doc.paragraphs)
for marker in ["Lê Đăng Tấn", "11236199", "CNTT65A", "TS. Lê Thị Hoài Thu", "tháng 12 năm 2026"]:
    assert marker in all_text, marker
assert len(doc.sections) == 3
assert len(doc.tables) >= 12
assert len(doc.inline_shapes) >= 10  # two logos, three diagrams, five design images
assert len(media) >= 9
assert "TOC " in xml and "SEQ Bảng" in xml and "SEQ Hình" in xml and "SEQ Công thức" in xml
normal = doc.styles["Normal"]
assert normal.font.name == "Times New Roman"
assert normal.font.size.pt == 13
assert normal.paragraph_format.line_spacing == 1.3
for section in doc.sections:
    assert round(section.left_margin.cm, 1) == 3.5
    assert round(section.right_margin.cm, 1) == 2.5
    assert round(section.top_margin.cm, 1) == 2.5
    assert round(section.bottom_margin.cm, 1) == 2.5
for label in ["Hình 3.1", "Hình 3.8", "Bảng 1.1", "Bảng 3.5", "Bảng 4.2"]:
    assert label in all_text, label
assert "Chưa có kết quả Turnitin" in all_text
print(f"PASS: {path.name}; {len(headings)} main headings; {len(doc.tables)} tables; "
      f"{len(doc.inline_shapes)} images; {len(doc.sections)} sections; {path.stat().st_size} bytes")
